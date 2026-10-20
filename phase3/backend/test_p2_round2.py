"""第四轮迭代回归测试（用户 2026-10-20 定的 4 项业务定义落地）。

运行：cd phase3 && ..\\.venv\\Scripts\\python.exe backend\\test_p2_round2.py
用临时数据库，不碰真实 pmc_phase3.db。
"""
import os, sys, tempfile, io
from datetime import date

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import backend.app as app_module
from fastapi.testclient import TestClient

PASS = 0
FAIL = []

def check(cond, msg):
    global PASS
    if cond:
        PASS += 1
        print("[OK] " + msg)
    else:
        FAIL.append(msg)
        print("[FAIL] " + msg)

def login(c, user):
    r = c.post("/api/auth/login", json={"username": user, "password": user})
    assert r.status_code == 200, r.text
    return {"Authorization": "Bearer " + r.json()["token"]}

def main():
    tmp = tempfile.NamedTemporaryFile(suffix=".db", delete=False)
    tmp.close()
    app_module.DB_PATH = tmp.name
    app_module.init_db()
    c = TestClient(app_module.app)
    admin = login(c, "admin")
    today_tag = "BATCH-" + date.today().strftime("%Y%m%d")

    # ---------- P1-5：仓库默认库位 ----------
    r = c.get("/api/warehouses", headers=admin)
    whs = {w["code"]: w.get("default_location") for w in r.json()}
    check(r.status_code == 200 and whs.get("主仓") == "A区-01", "P1-5 仓库主数据含默认库位，主仓=A区-01，实际 %s" % whs)
    check(whs.get("分仓") == "B区-01" and whs.get("成品仓") == "C区-01" and whs.get("上海仓") == "D区-01",
          "P1-5 分仓=B区-01、成品仓=C区-01、上海仓=D区-01，实际 %s" % whs)

    # 改主仓默认库位 → 入库不填库位应进新默认库位；再改回去
    r = c.put("/api/warehouses/主仓", headers=admin, json={"default_location": "测试-A"})
    check(r.status_code == 200, "P1-5 PUT 修改仓库默认库位，实际 %s %s" % (r.status_code, r.text[:120]))
    r = c.post("/api/stock/in", headers=admin, json={"part_code": "RM-001", "qty": 5, "warehouse": "主仓"})
    check(r.status_code == 200, "P1-5 入库不填库位/批次，实际 %s" % r.status_code)
    rows = [x for x in c.get("/api/stock", headers=admin).json() if x["part_code"] == "RM-001" and x["warehouse"] == "主仓"]
    hit = [x for x in rows if x["location"] == "测试-A" and float(x["qty"]) == 5.0]
    check(len(hit) == 1, "P1-5 不填库位→进了仓库默认库位(测试-A)，实际 %s" % [(x["location"], x["batch"], x["qty"]) for x in rows])
    check(len(hit) == 1 and hit[0]["batch"] == today_tag, "P1-5 不填批次→当天日期批次 %s，实际 %s" % (today_tag, hit[0]["batch"] if hit else None))
    c.put("/api/warehouses/主仓", headers=admin, json={"default_location": "A区-01"})
    r = c.get("/api/warehouses", headers=admin)
    check({w["code"]: w.get("default_location") for w in r.json()}.get("主仓") == "A区-01", "P1-5 默认库位改回去，主仓=A区-01")

    # 采购到货不填库位/批次 → 仓库默认库位 + 当天日期批次
    r = c.post("/api/purchase", headers=admin, json={"supplier_code": "SUP-001", "items": [{"part_code": "RM-001", "qty": 10}]})
    po_no = r.json().get("order_no")
    check(r.status_code == 200 and po_no, "P1-5 准备采购单，实际 %s %s" % (r.status_code, r.text[:120]))
    r = c.post("/api/purchase/receive", headers=admin, json={"order_no": po_no, "part_code": "RM-001", "qty": 10})
    check(r.status_code == 200, "P1-5 采购到货不填库位/批次，实际 %s %s" % (r.status_code, r.text[:120]))
    rows = [x for x in c.get("/api/stock", headers=admin).json() if x["part_code"] == "RM-001" and x["warehouse"] == "主仓"]
    hit = [x for x in rows if x["location"] == "A区-01" and x["batch"] == today_tag and float(x["qty"]) == 10.0]
    check(len(hit) == 1, "P1-5 到货进了默认库位A区-01+当天批次，实际 %s" % [(x["location"], x["batch"], x["qty"]) for x in rows])

    # ---------- P2-5：入库表/出库表导出 ----------
    r = c.get("/api/excel/export/stock_in", headers=admin)
    check(r.status_code == 200 and r.content[:2] == b"PK", "P2-5 入库表导出 200 且为真实 xlsx(PK 头)，实际 %s" % r.status_code)
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.content)); ws = wb.active
    headers = [v.value for v in ws[1]]
    exp_headers = ["单号", "日期", "类型", "仓库", "库位", "零件编码", "零件名称", "规格型号", "批次", "数量", "单价", "金额", "经办人", "关联单号", "备注"]
    check(headers == exp_headers, "P2-5 入库表 15 列表头与用户拟定一致，实际 %s" % headers)
    data = list(ws.iter_rows(min_row=2, values_only=True))
    po_rows = [x for x in data if x[13] == po_no and x[2] == "采购入库"]
    check(len(po_rows) >= 1 and str(po_rows[0][0]).startswith("RK-") and po_rows[0][12] == "admin",
          "P2-5 入库表含采购到货行：单号RK-开头、经办人admin，实际 %s" % (list(po_rows[0][:14]) if po_rows else None,))
    check(all(x[9] > 0 for x in data), "P2-5 入库表数量均为正数")

    r = c.post("/api/stock/out", headers=admin, json={"part_code": "RM-001", "qty": 3, "warehouse": "主仓"})
    check(r.status_code == 200, "P2-5 出库(不填批次走FIFO)，实际 %s %s" % (r.status_code, r.text[:120]))
    r = c.get("/api/excel/export/stock_out", headers=admin)
    wb = load_workbook(io.BytesIO(r.content)); ws = wb.active
    check(r.status_code == 200 and [v.value for v in ws[1]] == exp_headers, "P2-5 出库表导出 200 且 15 列表头一致")
    data = list(ws.iter_rows(min_row=2, values_only=True))
    out_rows = [x for x in data if x[5] == "RM-001" and x[2] == "其他出库"]
    check(len(out_rows) >= 1 and str(out_rows[0][0]).startswith("CK-") and float(out_rows[0][9]) == 3.0,
          "P2-5 出库表含出库行：单号CK-开头、数量3，实际 %s" % (list(out_rows[0][:14]) if out_rows else None,))

    # ---------- P2-6 4a：领料部门/用途（选填） ----------
    mo = c.post("/api/production", headers=admin, json={"product_code": "FG-001", "qty": 10}).json()["order_no"]
    r = c.post("/api/production/consume", headers=admin, json={"order_no": mo, "dept": "生产部", "purpose": "测试用途"})
    check(r.status_code == 200 and r.json().get("dept") == "生产部" and r.json().get("purpose") == "测试用途",
          "P2-6 领料带部门/用途，响应回带，实际 %s" % r.text[:160])
    rows = c.get(f"/api/production/{mo}/consumptions", headers=admin).json()
    check(len(rows) > 0 and all(x.get("dept") == "生产部" and x.get("purpose") == "测试用途" for x in rows),
          "P2-6 领料明细存下部门/用途，共 %d 条" % len(rows))
    # 不填部门/用途也能领（选填）；先把上一轮领空的 RM-003 补一点，避免缺料 400 干扰
    c.post("/api/stock/in", headers=admin, json={"part_code": "RM-003", "qty": 5, "warehouse": "主仓"})
    mo2 = c.post("/api/production", headers=admin, json={"product_code": "FG-001", "qty": 1}).json()["order_no"]
    r = c.post("/api/production/consume", headers=admin, json={"order_no": mo2})
    check(r.status_code == 200, "P2-6 部门/用途选填：不填也能领料，实际 %s %s" % (r.status_code, r.text[:120]))

    # ---------- P2-6 4b：在制不算可用库存 ----------
    r = c.get("/api/stock/available/FG-001", headers=admin)
    av = r.json()
    check(r.status_code == 200 and float(av.get("available", -1)) == 0.0,
          "P2-6 在制不算可用：领料后未完工，FG-001 可用量仍为 0，实际 %s" % av)
    r = c.post("/api/production/complete", headers=admin, json={"order_no": mo})
    check(r.status_code == 200, "P2-6 完工，实际 %s" % r.status_code)
    r = c.get("/api/stock/available/FG-001", headers=admin)
    check(float(r.json().get("available", -1)) == 10.0, "P2-6 完工后成品可用 10，实际 %s" % r.json())

    # ---------- P2-6 4c：齐套接口 + 看板 ----------
    mo3 = c.post("/api/production", headers=admin, json={"product_code": "FG-001", "qty": 100}).json()["order_no"]
    r = c.get(f"/api/production/{mo3}/kitting", headers=admin)
    kd = r.json()
    check(r.status_code == 200 and kd["total_items"] == 5, "P2-6 齐套接口：FG-001 直接子件 5 项，实际 %s" % kd.get("total_items"))
    check(kd["kitting_rate"] < 100 and any(not it["is_ready"] for it in kd["items"]),
          "P2-6 齐套率<100 且有未到齐项，实际到齐率 %s" % kd.get("kitting_rate"))
    check(all("shortage" in it and "available" in it for it in kd["items"]), "P2-6 齐套明细含可用/缺口")
    r = c.get("/api/dashboard", headers=admin)
    dash = r.json()
    check("kitting" in dash and dash["kitting"]["title"] == "齐套看板", "P2-6 看板第6块=齐套看板")
    kit_orders = [x["order_no"] for x in dash["kitting"]["items"]]
    check(mo3 in kit_orders, "P2-6 齐套看板含新生产单 %s" % mo3)
    rates = [x["kitting_rate"] for x in dash["kitting"]["items"]]
    check(rates == sorted(rates), "P2-6 齐套看板按到齐率从低到高排")

    # ---------- P1-9：保持现状（只提示不强制），此项关闭 ----------
    r = c.post("/api/auth/login", json={"username": "admin", "password": "admin"})
    check(r.status_code == 200 and r.json().get("must_change_password") is True, "P1-9 首次登录仍只提示改密(must_change_password)")
    r = c.get("/api/parts", headers=admin)
    check(r.status_code == 200, "P1-9 服务端不强制拦截：未改密仍可调接口(用户决定选B，此项关闭)")

    print()
    if FAIL:
        print("FAIL %d 项：" % len(FAIL))
        for f in FAIL:
            print("  - " + f)
        sys.exit(1)
    print("本轮 %d 项全部通过 ✅" % PASS)

if __name__ == "__main__":
    main()
